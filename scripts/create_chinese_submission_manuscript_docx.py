#!/usr/bin/env python3
"""Build a Chinese reading copy of the PlantMR submission manuscript."""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/manuscript-draft.zh-CN.md"
OUTPUT = ROOT / "docs/PlantMR_submission_manuscript.zh-CN.docx"
NAVY = "17365D"
TEAL = "0F6B78"
GREY = "F4F6F8"


def clean_inline(text):
    text = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    return text.strip()


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell(cell, text, bold=False, size=8.4, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(clean_inline(str(text)))
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, rows):
    if not rows:
        return
    headers = rows[0]
    body = [r for r in rows[1:] if not all(re.fullmatch(r"\s*:?-{3,}:?\s*", x) for x in r)]
    ncol = len(headers)
    table = doc.add_table(rows=1, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        set_cell(table.rows[0].cells[i], value, bold=True, size=8.2, color="FFFFFF")
        shade(table.rows[0].cells[i], NAVY)
    for ri, row in enumerate(body):
        cells = table.add_row().cells
        row = row + [""] * (ncol - len(row))
        for i in range(ncol):
            set_cell(cells[i], row[i], size=8.0)
            if ri % 2:
                shade(cells[i], GREY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page_field(paragraph):
    r = paragraph.add_run("第 ")
    r.font.size = Pt(8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    r._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    r._r.append(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    r._r.append(separate)
    value = OxmlElement("w:t")
    value.text = "1"
    r._r.append(value)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.append(end)
    r2 = paragraph.add_run(" 页")
    r2.font.size = Pt(8)


def is_caption(text):
    return bool(re.match(r"^(图|Figure)\s*\d+\.", text.strip()))


def build():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for level, size, color in [(1, 16, NAVY), (2, 13, TEAL), (3, 11, TEAL)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)
    header = sec.header.paragraphs[0]
    header.text = "PlantMR｜中文阅读稿"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("666666")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_field(footer)

    title_done = False
    figures = 0
    tables = 0
    references = 0
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# ") and not title_done:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(30)
            r = p.add_run(clean_inline(line[2:]))
            r.bold = True
            r.font.size = Pt(19)
            r.font.color.rgb = RGBColor.from_string(NAVY)
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.add_run("中文阅读稿｜英文投稿稿的完整翻译，统计结果和引用保持一致").italic = True
            p3 = doc.add_paragraph()
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p3.add_run("PlantMR v1.1.0-paper").font.size = Pt(10)
            doc.add_page_break()
            title_done = True
            i += 1
            continue
        match = re.match(r"^(#{1,3})\s+(.*)$", line)
        if match:
            level = len(match.group(1))
            text = clean_inline(match.group(2))
            doc.add_paragraph(text, style=f"Heading {level}")
            if level == 1 and text in ("参考文献", "References"):
                pass
            i += 1
            continue
        if line.startswith("!["):
            match = re.match(r"!\[[^]]*\]\(([^)]+)\)", line)
            if match:
                rel = match.group(1)
                image = ROOT / "docs" / rel if rel.startswith("figures/") else ROOT / rel
                if image.exists():
                    doc.add_picture(str(image), width=Inches(6.35))
                    figures += 1
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                rows.append([x.strip() for x in raw.split("|")])
                i += 1
            add_table(doc, rows)
            tables += 1
            continue
        if line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            p.add_run(clean_inline(line[2:]))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            references += 1 if line[: line.find(".")].isdigit() else 0
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(clean_inline(line))
        if is_caption(line):
            r.italic = True
            r.font.size = Pt(9)
        i += 1

    doc.save(OUTPUT)
    print(OUTPUT)
    print(f"figures={figures} tables={tables} references={references}")


if __name__ == "__main__":
    build()
