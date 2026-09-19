#!/usr/bin/env python3
"""Create the Chinese PlantMR research explanation Word document."""
from pathlib import Path
import csv
import json
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "PlantMR研究内容与投稿准备说明.docx"
FIGURE = ROOT / "docs" / "figures" / "gxe_simulation.png"

NAVY = "17365D"
TEAL = "0F6B78"
LIGHT = "EAF3F5"
GREY = "F2F2F2"
DARK_GREY = "666666"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:%s" % edge
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in kwargs[edge]:
                    element.set(qn("w:%s" % key), str(kwargs[edge][key]))


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, value in enumerate(headers):
        set_cell_text(hdr[i], value, bold=True, color="FFFFFF", size=font_size)
        set_cell_shading(hdr[i], NAVY)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
            if row_index % 2 == 1:
                set_cell_shading(cells[i], "F7FAFB")
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_hyperlink(paragraph, text, url, color="0563C1", underline=True):
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    run.font.size = Pt(9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run2 = paragraph.add_run(" 页")
    run2.font.size = Pt(9)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p


def add_note(doc, title, text, fill=LIGHT):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, top={"val": "single", "sz": 8, "color": TEAL}, bottom={"val": "single", "sz": 8, "color": TEAL}, left={"val": "single", "sz": 8, "color": TEAL}, right={"val": "single", "sz": 8, "color": TEAL})
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    r.font.size = Pt(10)
    r2 = p.add_run(text)
    r2.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style="Heading %d" % level)
    p.add_run(text)
    return p


def normal(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def read_tsv(path):
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.1)
    sec.right_margin = Cm(2.1)

    styles = doc.styles
    normal_style = styles["Normal"]
    normal_style.font.name = "Microsoft YaHei"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal_style.font.size = Pt(10.5)
    for level, size, color in [(1, 16, NAVY), (2, 13, TEAL), (3, 11, TEAL)]:
        style = styles["Heading %d" % level]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if level == 1 else 7)
        style.paragraph_format.space_after = Pt(4)
    for style_name in ["List Bullet", "List Bullet 2", "List Number"]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(10)

    # Header/footer
    header = sec.header.paragraphs[0]
    header.text = "PlantMR 研究内容与投稿准备说明"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(DARK_GREY)
    footer = sec.footer.paragraphs[0]
    footer.text = "PlantMR v1.1.0-paper  |  "
    for run in footer.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(DARK_GREY)
    add_page_number(footer)

    # Title page
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(55)
    r = p.add_run("PlantMR研究内容与投稿准备说明")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("植物/作物摘要统计孟德尔随机化与环境效应异质性分析")
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(25)
    r = p3.add_run("研究版本：PlantMR v1.1.0\n论文包：v1.1.0-paper\n整理日期：2026年9月19日")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(DARK_GREY)
    add_note(doc, "一句话概括", "我们不是简单把人类MR代码换成植物数据，而是建立了一个明确声明物种、参考组装、环境、倍性、LD和样本重叠边界的植物MR工作流，并增加了一个能处理环境效应异质性的协方差感知模型。", fill="E8F2F7")
    doc.add_page_break()

    # Contents (manual stable TOC)
    heading(doc, "目录", 1)
    toc = [
        "一、当前结论和研究定位",
        "二、为什么要做植物版MR",
        "三、与已有方法和软件的关系",
        "四、我们实际实现了什么",
        "五、核心统计模型",
        "六、模拟验证设计与结果",
        "七、Arabidopsis真实数据案例",
        "八、软件结构、输入输出与复现",
        "九、当前不能声称的内容",
        "十、投稿期刊建议",
        "十一、下一步投稿前Gate",
        "十二、参考资料与项目文件",
    ]
    for i in toc:
        add_bullet(doc, i)
    doc.add_page_break()

    heading(doc, "一、当前结论和研究定位", 1)
    normal(doc, "当前研究已经从“植物MR软件原型”推进到“可进行编辑预询和方法学预审的软件/方法论文包”。软件、模拟、真实植物数据案例、论文初稿、审稿风险表和投稿清单均已生成并通过本机验证。")
    add_table(doc, ["项目项", "当前状态", "说明"], [
        ["软件版本", "v1.1.0", "Python CLI与库；MIT许可证"],
        ["论文包", "v1.1.0-paper", "包含论文、模拟、真实案例、图、表和投稿清单"],
        ["自动化测试", "23个通过", "包括GxE完整网格和奇异LD自由度测试"],
        ["主要方法", "环境效应异质性GLS", "不是MR-GxE多效性校正的替代品"],
        ["真实案例", "Arabidopsis AT1G11560", "数据契约和审计案例，不是独立因果验证"],
        ["当前首选期刊", "Plant Methods", "最符合软件/植物方法论文定位"],
    ], [3.0, 3.8, 9.0])
    add_note(doc, "最重要的边界", "当前稿件适合投软件/方法方向，不适合包装成“发现AT1G11560因果作用”的生物学应用论文。Arabidopsis案例用于证明软件能够处理真实植物数据、LD和环境分层，不作为独立验证。", fill="FFF2CC")

    heading(doc, "二、为什么要做植物版MR", 1)
    normal(doc, "孟德尔随机化（Mendelian randomization，MR）利用遗传变异作为工具变量，估计遗传预测的分子暴露对表型结局的影响。人类MR的常规流程通常默认研究群体、参考基因组、LD面板、样本重叠和暴露定义比较标准化；植物数据往往不是这样。")
    add_bullet(doc, "植物自然群体常有明显地理结构、育种历史和亲缘关系，遗传关联容易混入群体结构。")
    add_bullet(doc, "不同植物群体的LD衰减速度差异较大，局部eQTL常包含高度相关的SNP，不能把每个SNP当作独立工具。")
    add_bullet(doc, "表达数据高度依赖组织、发育阶段、处理条件和环境，WW/DS、10°C/16°C等条件不能简单合并。")
    add_bullet(doc, "植物还可能有多倍体剂量、PAV、SV、同源亚基因组和泛基因组坐标问题，不能静默按普通二等位SNP处理。")
    normal(doc, "因此，植物MR软件的第一任务不是增加更多MR估计器，而是把输入数据的生物学语境和统计假设写进机器可检查的分析契约。")

    heading(doc, "三、与已有方法和软件的关系", 1)
    normal(doc, "我们专门查阅了植物MR、MR-GxE和异质GWAS/MR方法。结论是：环境交互MR并不是空白领域，不能把“植物版实现”写成理论首创。")
    add_table(doc, ["方法/软件", "主要目标", "与PlantMR的关系", "当前处理"], [
        ["MRBIGR", "植物多组学、GWAS、MR、网络和富集分析工作流", "植物软件方向的重要先例", "作为植物软件比较对象"],
        ["原玉米干旱MR", "使用lead-eQTL和2SLS优先候选基因", "提供植物MR应用基线", "用于数据和方法审计"],
        ["MR-GxE", "利用基因×协变量交互识别/校正多效性", "目标和截距解释不同", "不静默替代"],
        ["MR-GENIUS", "不要求显式观测交互，但依赖全局假设", "统计识别假设不同", "未内置"],
        ["MR-EILLS", "跨异质GWAS数据估计不变因果效应", "不是同一环境效应异质性估计量", "列为后续比较器"],
        ["PlantMR环境斜率", "估计MR效应随预设环境分数变化的斜率", "本文的核心实现", "已实现并验证"],
    ], [3.0, 5.0, 5.0, 3.0], font_size=8)
    add_note(doc, "定位变化", "PlantMR的可发表贡献不是“提出环境交互MR”，而是把植物数据契约、环境/LD协方差、完整SNP×环境网格和可复现审计组合成一个能实际运行的植物MR工具。", fill="E8F2F7")

    heading(doc, "四、我们实际实现了什么", 1)
    heading(doc, "4.1 普通摘要统计MR", 2)
    add_bullet(doc, "输入暴露和结局摘要统计，要求SNP、效应等位基因、另一等位基因、beta、SE和P值。")
    add_bullet(doc, "进行等位基因协调，包括反向等位基因、互补链和回文位点处理。")
    add_bullet(doc, "根据暴露P值、MAF和F统计量筛选工具变量。")
    add_bullet(doc, "支持Wald ratio、固定效应IVW、随机效应IVW、MR-Egger、Cochran异质性和逐工具变量留一分析。")
    add_bullet(doc, "可输入带符号LD相关矩阵进行LD clumping，并在缺失LD信息时显式警告。")
    heading(doc, "4.2 环境分层分析", 2)
    normal(doc, "`run-stratified`在每个环境中独立运行MR，输出每个环境的固定/随机IVW和MR-Egger结果。它只是条件性分层，不等于正式环境效应异质性模型。")
    heading(doc, "4.3 环境效应异质性模型", 2)
    normal(doc, "`run-gxe`要求每个SNP在所有环境中都有记录，并只保留在所有环境均通过P值、F统计量和MAF筛选的SNP。它对每个SNP×环境单元计算比值估计，再用两参数GLS估计总体效应和环境斜率。")
    add_table(doc, ["输出量", "含义", "不能解释成"], [
        ["theta_0 / intercept", "环境分数z=0时的遗传预测暴露效应", "MR-GxE多效性截距"],
        ["theta_1 / slope", "环境分数每增加一个单位时，MR效应的变化", "环境处理本身的因果效应"],
        ["Q和Q P值", "模型残差异质性诊断", "因果关系已被证实"],
        ["covariance_rank", "LD/环境协方差矩阵的数值秩", "工具变量有效性证明"],
    ], [4.0, 7.0, 5.0], font_size=8.5)

    heading(doc, "五、核心统计模型", 1)
    normal(doc, "对第j个SNP和第k个环境，先计算比值估计：")
    add_note(doc, "比值估计", "r_jk = beta_y,jk / beta_x,jk\n\n其中beta_x是SNP对分子暴露的效应，beta_y是SNP对表型结局的效应。", fill="F2F2F2")
    normal(doc, "环境效应模型为：")
    add_note(doc, "两参数模型", "r_jk = theta_0 + theta_1 z_k + epsilon_jk\n\n这里z_k是预先指定的环境数值。它可以是温度、干旱等级、处理强度或经过中心化的环境编码。工具不会自动替用户选择环境编码。", fill="F2F2F2")
    normal(doc, "比值估计的近似方差由暴露和结局两部分组成：")
    add_note(doc, "一阶delta方差", "v_jk = se_y,jk² / beta_x,jk² + beta_y,jk² × se_x,jk² / beta_x,jk⁴", fill="F2F2F2")
    normal(doc, "如果提供SNP-LD相关矩阵R_SNP和环境误差相关矩阵R_ENV，工具构造：")
    add_note(doc, "协方差矩阵", "V = D (R_SNP ⊗ R_ENV) D\n\nD是由每个比值标准误组成的对角矩阵，⊗表示Kronecker积。缺失的相关矩阵按单位矩阵处理，但报告会明确警告。", fill="F2F2F2")
    normal(doc, "GLS估计为：")
    add_note(doc, "GLS估计", "theta_hat = (X'V⁻¹X)⁻¹ X'V⁻¹r\n\nX的两列为[1,z]。如果LD导致V秩亏，Q检验自由度使用rank(V)-rank(X)，避免把高度相关的SNP重复计数。", fill="F2F2F2")
    add_note(doc, "统计边界", "该模型是环境效应异质性模型，不是MR-GxE多效性校正模型。它不能自动消除水平多效性，也不能从边际摘要统计中自动恢复样本重叠协方差。", fill="FFF2CC")

    heading(doc, "六、模拟验证设计与结果", 1)
    normal(doc, "我们把模拟结果和真实植物数据严格分开。模拟只验证代码和统计实现，不作为植物生物学证据。")
    add_table(doc, ["模拟组", "设置", "主要结果"], [
        ["Null", "20 SNP、4环境、环境相关0.5、500次", "正确模型拒绝率0.054，覆盖率0.946"],
        ["Causal G×E", "真实环境斜率0.25、500次", "均值0.24982，偏差−0.00018，功效1.00"],
        ["Directional pleiotropy", "共同多效性0.04、真实斜率0", "斜率近似无偏，但截距偏差+0.216"],
        ["LD stress", "AR(1) LD rho=0.6、环境相关0.5、500次", "正确模型拒绝率0.044；独立性错配0.126"],
    ], [4.0, 7.0, 6.0], font_size=8.5)
    normal(doc, "下图是主要模拟图。左侧展示环境斜率的抽样分布，右侧展示拒绝率和置信区间覆盖率。Null场景的拒绝率应接近0.05，Causal G×E场景的功效应接近1，方向性多效性场景用于展示模型的失效边界。")
    if FIGURE.exists():
        doc.add_picture(str(FIGURE), width=Inches(6.35))
        cap = doc.add_paragraph("图1  PlantMR环境斜率模拟结果。蓝色为协方差感知模型，橙色为对角协方差比较模型。该图只代表模拟验证，不代表任何真实植物基因的因果结果。")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(9)
    normal(doc, "LD压力测试是本轮新增的关键增强。真实错误结构同时包含环境相关和SNP间LD相关；如果按独立SNP处理，零效应下的错误拒绝率从0.044升到0.126，95%覆盖率从0.956降到0.874。这个结果支持把LD协方差作为主要分析输入，而不是可选的装饰信息。")

    heading(doc, "七、Arabidopsis真实数据案例", 1)
    heading(doc, "7.1 案例目的", 2)
    normal(doc, "这个案例不是为了重新发现AT1G11560，而是为了验证PlantMR能否在真实植物数据上完成：数据下载、ID对应、等位基因映射、群体结构校正、局部LD、环境标签、完整网格筛选和可审计报告。")
    heading(doc, "7.2 数据组成", 2)
    add_table(doc, ["数据层", "来源", "本项目处理"], [
        ["分子暴露", "GSE80744归一化表达矩阵", "提取AT1G11560，log1p处理，作为基线暴露"],
        ["基因型", "1001 Genomes v3.1", "提取Chr1:3,861,124–3,901,085局部区域，共3,352个SNP"],
        ["群体结构", "同一1001 Genomes矩阵", "从稀疏全基因组标记计算PC1–PC5"],
        ["结局环境1", "AraPheno FT10", "10°C开花时间，环境分数−3"],
        ["结局环境2", "AraPheno FT16", "16°C开花时间，环境分数+3"],
        ["等位基因", "1001 Genomes accession VCF", "用于把二进制剂量映射为真实REF/ALT"],
    ], [3.2, 5.0, 8.0], font_size=8.2)
    heading(doc, "7.3 工具筛选和主要结果", 2)
    add_bullet(doc, "预设工具条件：暴露P≤5×10^-8、F≥10、MAF≥0.05。")
    add_bullet(doc, "要求SNP在两个环境都通过筛选，最终保留48个工具变量。")
    add_bullet(doc, "48个工具高度相关，因此主分析输入签名LD矩阵。")
    add_bullet(doc, "主分析环境斜率−0.03705，SE=0.15617，P=0.81249。")
    add_bullet(doc, "主分析截距2.3601，SE=0.4685，P=4.72×10^-7。")
    add_bullet(doc, "残余异质性Q=192.89，秩调整后df=60，P=7.33×10^-16。")
    add_bullet(doc, "用FT10/FT16表型相关r=0.88195做敏感性代理时，斜率−0.08228，P=0.21491；这不是已知的ratio-error covariance。")
    add_note(doc, "真实案例结论", "当前数据不支持AT1G11560在10°C和16°C之间存在显著MR效应差异。更重要的发现是：局部LD和环境协方差会实质改变不确定性，因此不能把独立SNP的普通IVW当作默认答案。", fill="E8F2F7")
    heading(doc, "7.4 为什么不能称为独立验证", 2)
    add_bullet(doc, "表达和FT10/FT16数据共享accessions，样本重叠协方差没有被精确估计。")
    add_bullet(doc, "本案例使用local OLS+PC1–PC5，不是原论文的正式LMM/SMR结果。")
    add_bullet(doc, "基线表达暴露复制到两个结局环境，斜率表示结局环境上的效应异质性，不是环境特异eQTL机制。")
    add_bullet(doc, "没有把已有论文的SMR/HEIDI结果伪装成PlantMR独立验证。")

    heading(doc, "八、软件结构、输入输出与复现", 1)
    add_table(doc, ["路径", "作用"], [
        ["src/plant_mr/schema.py", "摘要统计输入验证和植物元数据"],
        ["src/plant_mr/harmonize.py", "等位基因协调和回文位点处理"],
        ["src/plant_mr/instruments.py", "P值、MAF、F统计量和LD筛选"],
        ["src/plant_mr/estimators.py", "Wald、IVW、随机效应、MR-Egger和留一分析"],
        ["src/plant_mr/gxe.py", "环境斜率GLS、协方差和秩调整Q检验"],
        ["src/plant_mr/cli.py", "validate、run、run-stratified、run-gxe命令"],
        ["scripts/run_gxe_simulation.py", "基础模拟"],
        ["scripts/run_gxe_ld_stress.py", "LD/环境相关压力测试"],
        ["scripts/benchmark_runtime.py", "运行时间和峰值RSS"],
    ], [5.5, 11.0], font_size=8.4)
    heading(doc, "8.1 最小复现命令", 2)
    for cmd in [
        "python -m pip install --no-deps -e .",
        "plantmr validate --exposure examples/synthetic/exposure.tsv --outcome examples/synthetic/outcome.tsv --metadata examples/synthetic/metadata.json",
        "plantmr run-gxe --exposure examples/synthetic/environment_exposure.tsv --outcome examples/synthetic/environment_outcome.tsv --metadata examples/synthetic/environment_metadata.json --outdir gxe_result",
        "pytest -q",
        "python scripts/run_gxe_simulation.py --reps 500 --outdir results/benchmarks/gxe_simulation",
        "python scripts/run_gxe_ld_stress.py --reps 500 --outdir results/benchmarks/gxe_ld_stress",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(cmd)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
    normal(doc, "验证主机上，普通验证约9.7秒，合成GxE约9.8秒，Arabidopsis真实案例约9.9秒，峰值RSS约123–131MB。容器构建在当前主机未完成，因为没有可用Docker/Podman运行时。")

    heading(doc, "九、当前不能声称的内容", 1)
    add_table(doc, ["不能这样写", "应该这样写"], [
        ["PlantMR首次提出环境交互MR", "PlantMR提供植物数据契约和环境效应异质性实现；已有MR-GxE等相关方法"],
        ["AT1G11560已被证明为因果基因", "AT1G11560案例用于展示真实数据契约和审计流程，不能作为独立因果验证"],
        ["PlantMR消除了水平多效性", "方向性多效性模拟显示截距仍会偏差，模型不自动修复多效性"],
        ["环境斜率就是环境处理的因果效应", "环境斜率表示预设环境分数下遗传预测暴露效应的变化"],
        ["玉米97个候选基因是本研究新发现", "玉米补充表用于原研究候选清单和数据链审计"],
        ["独立IVW结果就是主结果", "局部LD下独立IVW仅作为比较，主结果使用LD协方差"],
    ], [7.0, 10.0], font_size=8.5)

    heading(doc, "十、投稿期刊建议", 1)
    add_table(doc, ["期刊", "匹配度", "建议"], [
        ["Plant Methods", "高", "首选；先做预投稿咨询，完成公开仓库/DOI后正式投稿"],
        ["BMC Bioinformatics", "中高", "补MR-GxE/MR-GENIUS/MR-EILLS实现级比较后考虑"],
        ["GigaScience", "中高", "把软件、数据、模拟、环境锁和结果做成FAIR研究对象"],
        ["Plant Communications", "中", "需要更强的植物生物学问题和独立验证"],
        ["PLOS Computational Biology", "目前偏低", "需要更强算法创新、广泛比较和生物学洞见"],
        ["Bioinformatics", "目前偏低", "需要更强算法/规模化基准和直接工具比较"],
        ["Nature Methods/Communications Biology", "暂不建议", "当前创新和独立生物学验证不足"],
    ], [4.0, 3.0, 10.0], font_size=8.2)
    normal(doc, "当前最稳妥的路线是：先向Plant Methods做预投稿咨询；完成公开代码仓库、Zenodo DOI、干净环境复现和作者信息后正式投稿。")

    heading(doc, "十一、下一步投稿前Gate", 1)
    add_number(doc, "把本地仓库推送到公开GitHub或GitLab，保持v1.1.0和v1.1.0-paper标签不变。")
    add_number(doc, "通过Zenodo或等价服务生成DOI，并把DOI写入论文Code availability。")
    add_number(doc, "在干净Python环境或容器中重跑23个测试、合成demo、基础模拟和LD压力测试。")
    add_number(doc, "邀请不了解项目的合作者独立安装和运行demo，记录遇到的问题。")
    add_number(doc, "补齐作者、基金、贡献、利益冲突、通讯作者和目标期刊格式。")
    add_number(doc, "如果改投生物学应用论文，再补正式LMM摘要统计、样本重叠协方差、独立验证、共定位/HEIDI和功能实验。")
    add_note(doc, "投稿安全线", "软件/方法论文可以在公开仓库和独立复现补齐后投稿；生物学因果论文不能仅凭当前Arabidopsis案例投稿。", fill="FFF2CC")

    heading(doc, "十二、参考资料与项目文件", 1)
    refs = [
        ("MRBIGR植物多组学MR工具", "https://doi.org/10.1016/j.xplc.2024.101197"),
        ("玉米干旱eQTL与MR研究", "https://doi.org/10.1186/s13059-020-02069-1"),
        ("MR-GxE", "https://doi.org/10.1093/ije/dyy204"),
        ("MR-GENIUS/interaction-based MR", "https://doi.org/10.1371/journal.pone.0271933"),
        ("MR-EILLS", "https://doi.org/10.1038/s41467-025-62823-6"),
        ("Arabidopsis双性状和SMR/HEIDI案例", "https://doi.org/10.1038/s41437-024-00688-z"),
        ("1001 Genomes", "https://doi.org/10.1016/j.cell.2016.05.063"),
        ("STROBE-MR", "https://pubmed.ncbi.nlm.nih.gov/34698778/"),
        ("Plant Methods范围", "https://plantmethods.biomedcentral.com/about"),
        ("BMC Bioinformatics软件文章要求", "https://link.springer.com/journal/12859/submission-guidelines/software-article"),
        ("GigaScience开放科学范围", "https://academic.oup.com/gigascience/pages/About"),
        ("Nature代码/软件指南", "https://media.nature.com/full/nature-cms/documents/GuidelinesCodePublication.pdf"),
    ]
    for i, (name, url) in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.add_run("[%d] %s：" % (i, name))
        add_hyperlink(p, url, url)
    normal(doc, "项目核心文件：docs/manuscript-draft.en.md、docs/journal-recommendation.zh-CN.md、docs/reviewer-risk-audit.zh-CN.md、docs/software-submission-checklist.md、docs/strobe-mr-crosswalk.md。")

    # Core properties
    props = doc.core_properties
    props.title = "PlantMR研究内容与投稿准备说明"
    props.subject = "植物/作物摘要统计MR与环境效应异质性"
    props.author = "PlantMR contributors"
    props.keywords = "PlantMR, Mendelian randomization, GxE, LD, plant methods"
    props.comments = "Generated from the verified PlantMR v1.1.0-paper research package."

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
