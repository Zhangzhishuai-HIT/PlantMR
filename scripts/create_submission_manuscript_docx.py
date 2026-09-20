#!/usr/bin/env python3
"""Create the full English submission manuscript as a DOCX."""
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/PlantMR_submission_manuscript.docx'
FIG=ROOT/'docs/figures/gxe_simulation.png'
FIG_WORKFLOW=ROOT/'docs/figures/plantmr_workflow.png'
FIG_LD=ROOT/'docs/figures/gxe_ld_stress.png'
FIG_CASE=ROOT/'docs/figures/arabidopsis_case.png'
NAVY='17365D'; TEAL='0F6B78'; LIGHT='EAF3F5'; GREY='F4F6F8'; YELLOW='FFF2CC'


def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def cell_text(cell,text,bold=False,size=8.5,color=None):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc,headers,rows,widths=None,size=8.5):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers): cell_text(t.rows[0].cells[i],h,True,size,'FFFFFF'); shade(t.rows[0].cells[i],NAVY)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cell_text(cells[i],v,size=size)
            if ri%2: shade(cells[i],GREY)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def heading(doc,text,level=1):
    p=doc.add_paragraph(style='Heading %d'%level); p.add_run(text); return p

def para(doc,text,after=5):
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.12; p.paragraph_format.space_after=Pt(after); p.add_run(text); return p

def bullet(doc,text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(text); return p

def numbered(doc,text):
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(2); p.add_run(text); return p

def note(doc,title,text,fill=LIGHT):
    t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,fill); c.text=''
    p=c.paragraphs[0]; r=p.add_run(title+'\n'); r.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL); r.font.size=Pt(9.5)
    r=p.add_run(text); r.font.size=Pt(9)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def hyperlink(paragraph,text,url):
    rid=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid); r=OxmlElement('w:r'); rp=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),'0563C1'); rp.append(c); u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rp.append(u); r.append(rp); t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); paragraph._p.append(h)

def page_field(p):
    r=p.add_run('Page '); r.font.size=Pt(8)
    for typ in ['begin','separate','end']:
        el=OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'),typ); r._r.append(el)
        if typ=='separate':
            t=OxmlElement('w:t'); t.text='1'; r._r.append(t)

def main():
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2.0); sec.right_margin=Cm(2.0)
    styles=doc.styles; n=styles['Normal']; n.font.name='Arial'; n._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); n.font.size=Pt(10.5)
    for level,size,color in [(1,16,NAVY),(2,13,TEAL),(3,11,TEAL)]:
        s=styles['Heading %d'%level]; s.font.name='Arial'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(10); s.paragraph_format.space_after=Pt(4)
    for sn in ['List Bullet','List Number']:
        s=styles[sn]; s.font.name='Arial'; s.font.size=Pt(10)
    header=sec.header.paragraphs[0]; header.text='PlantMR | Submission manuscript'; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string('666666')
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; page_field(footer)

    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(35)
    r=p.add_run('PlantMR: an auditable summary-statistics toolkit for plant and crop Mendelian randomization with environment-aware effect heterogeneity'); r.bold=True; r.font.size=Pt(20); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('[Author names, affiliations and corresponding-author details to be supplied]').italic=True
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Submission manuscript | PlantMR v1.1.0-paper | 19 September 2026').font.size=Pt(10)
    note(doc,'Submission status','The scientific content, analyses, source data, simulations and software snapshot are frozen. Author metadata, public repository DOI and final journal formatting remain to be supplied before external submission.',YELLOW)
    doc.add_page_break()

    heading(doc,'Abstract',1)
    heading(doc,'Background',2)
    para(doc,'Plant Mendelian randomization (MR) can connect molecular traits to agronomic phenotypes, but plant reference assemblies, environmental strata, linkage disequilibrium (LD) and sample overlap are often implicit. We developed PlantMR 1.1 as a local, open-source toolkit for auditable plant and crop summary-statistics MR with a narrowly defined environment-effect-heterogeneity analysis.')
    heading(doc,'Results',2)
    para(doc,'PlantMR validates plant metadata and summary statistics, harmonizes alleles, enforces complete SNP-by-environment instrument grids, accepts signed SNP-LD and environment-correlation matrices, and reports rank-aware heterogeneity diagnostics. In 500-replicate simulations, the covariance-aware estimator had a null rejection rate of 0.054 and 95% coverage of 0.946; for a true environment slope of 0.25, the mean estimate was 0.2498, bias −0.00018, coverage 0.952 and power 1.00. In an LD stress test with AR(1) LD correlation ρ=0.6, correct covariance modeling gave rejection 0.044 and coverage 0.956, whereas an independence approximation gave 0.126 and 0.874. A real Arabidopsis data-contract case retained 48 LD-correlated instruments for AT1G11560 and flowering time at 10°C and 16°C. The primary slope was −0.0371 (SE 0.1562, P=0.812), with residual heterogeneity Q=192.89, df=60, P=7.33×10−16.')
    heading(doc,'Conclusions',2)
    para(doc,'PlantMR provides a reproducible plant-focused implementation and audit workflow. The simulations support covariance-aware calibration under the tested scenarios, but the Arabidopsis case is a software demonstration rather than independent causal validation. PlantMR does not replace colocalization, formal interaction-MR methods, mixed-model GWAS or functional validation.')
    para(doc,'Keywords: Mendelian randomization; plant genomics; crop genomics; genotype–environment interaction; eQTL; linkage disequilibrium; summary statistics; reproducibility')

    heading(doc,'Background',1)
    para(doc,'Mendelian randomization uses genetic associations as instrumental variables to estimate genetically proxied exposure effects on outcomes.[1,2] The relevance, exchangeability and exclusion-restriction assumptions provide the foundation for interpretation, but their implementation depends on the study population, phenotype, molecular exposure and data-generating design. Summary-statistics MR has enabled large-scale analyses using public genome-wide association studies (GWAS) and platforms such as MR-Base,[10] yet summarized associations are also vulnerable to weak instruments, horizontal pleiotropy, LD and sample overlap.[3–9]')
    para(doc,'Plant applications add several complications. Natural accessions and breeding panels can exhibit strong geographic structure and kinship. LD decay can differ substantially between species, populations and genomic regions. Expression is strongly dependent on tissue, developmental stage and treatment. Multi-environment trials and stress experiments frequently measure the same genotypes under several conditions, creating correlated estimates rather than independent GWAS datasets. Polyploid dosage, presence/absence variation and structural variation introduce additional representations that cannot be silently converted to ordinary biallelic SNPs; recent Arabidopsis pan-genome work illustrates why local adaptation and non-reference sequence variation should remain explicit.[35]')
    para(doc,'Existing plant-oriented software demonstrates the value of integrated workflows. MRBIGR provides a broad population-scale multi-omics and MR toolbox for maize.[17] The maize drought study that motivates the present work prioritized genes using dynamic eQTLs and a lead-eQTL MR workflow.[18] Plant G×E GWAS and meta-analysis methods also show that environmental contrasts, population structure and interaction terms must be modeled jointly rather than treated as labels.[25–27] Outside plants, MR-GxE uses gene-by-covariate interactions to detect or correct pleiotropic bias under a constant-pleiotropy assumption,[14] whereas interaction-based MR work distinguishes MR-GxE from MR-GENIUS and emphasizes interaction strength and assumption sensitivity.[15] MR-EILLS targets an invariant causal effect across heterogeneous GWAS summary datasets and evaluates multiple MR and meta-analytic comparators.[16] These methods have different estimands and should not be conflated.')
    para(doc,'Here we develop PlantMR as a plant/crop summary-statistics software and reporting workflow. The contribution is deliberately narrower than a new general theory of interaction MR: (i) an explicit plant metadata and summary-statistics contract; (ii) complete SNP-by-environment instrument selection; (iii) signed LD and environment-correlation inputs; (iv) rank-aware heterogeneity diagnostics; and (v) a reproducible local implementation with simulations and a real Arabidopsis data-contract case. We evaluate calibration, LD-misspecification risk and failure boundaries rather than claiming a new causal gene discovery.')

    heading(doc,'Implementation',1)
    heading(doc,'Study questions and workflow',2)
    para(doc,'We evaluated four prespecified questions: (1) can the software enforce a plant-aware, auditable data contract; (2) does the covariance-aware environment-slope estimator recover known effects under correlated SNP and environment errors; (3) how much can inference deteriorate when LD and environment dependence are incorrectly treated as independent; and (4) can the same audit trail be executed on a real plant data contract without converting a demonstration into a causal-gene claim? The workflow in Figure 1 separates these questions into input provenance, harmonization, instrument selection, covariance specification, estimation and reporting.')
    if FIG_WORKFLOW.exists():
        doc.add_picture(str(FIG_WORKFLOW),width=Inches(6.55))
        p=doc.add_paragraph('Figure 1. PlantMR reproducible workflow. Plant metadata and summary statistics are retained alongside allele harmonization, complete-grid quality control, covariance specification and machine-readable audit outputs.')
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic=True; p.runs[0].font.size=Pt(9)
    heading(doc,'PlantMR software contract and analysis workflow',2)
    para(doc,'PlantMR accepts exposure and outcome tables containing SNP, effect allele, other allele, beta, standard error and P value. Optional fields include effect-allele frequency and sample size. Plant metadata records species, reference assembly, trait, tissue, developmental stage, environment, ploidy and LD-panel provenance. The software validates numeric values and alleles, harmonizes reversed and complemented alleles, identifies ambiguous palindromic variants, applies P-value/MAF/F-statistic filters and records every exclusion.')
    para(doc,'Table 1. PlantMR components, boundaries and audit outputs.',after=2)
    table(doc,['Component','Function','Output'],[
        ['Schema and metadata','Validate summary statistics and plant context','Validation messages and metadata JSON'],
        ['Allele harmonization','Align effect alleles and handle palindromic variants','Harmonized table and exclusion counts'],
        ['Instrument QC','P value, MAF, F statistic and complete-grid filtering','Instrument audit and retained SNP table'],
        ['Standard MR','Wald, fixed/random IVW, MR-Egger, Q and leave-one-out','JSON, TSV and Markdown reports'],
        ['Environment slope','GLS on SNP×environment ratio estimates','Pooled effect, slope, covariance rank and Q'],
        ['Reproducibility','CLI, tests, simulations, source-data records','Versioned reports and benchmark outputs'],
    ],[3.2,7.0,6.0],8.2)
    para(doc,'The command-line interface exposes validate, run, run-stratified and run-gxe commands. The software does not silently substitute SMR, GSMR, colocalization, MVMR, MR-PRESSO, MR-GxE, MR-GENIUS or MR-EILLS.')

    heading(doc,'Comparison with related tools',2)
    para(doc,'Following the structure used in recent Plant Methods software articles, we compare PlantMR with functionally adjacent tools before presenting the benchmarks. This is a scope and contract comparison based on the cited software papers and public documentation, not a fabricated head-to-head runtime experiment. PlantMR’s distinct target is the combination of plant metadata, complete SNP-by-environment grids, signed LD/environment covariance and auditable summary-statistics outputs.')
    para(doc,'Table 2. Functional positioning against related tools. “Not the primary focus” means that a tool may support a related analysis but does not expose the same contract or estimand.',after=2)
    table(doc,['Tool','Primary scope','Plant context','Environment / LD handling','Evidence used here'],[
        ['PlantMR 1.1','Plant summary-statistics MR and environment-effect heterogeneity','Native metadata, assembly, tissue, stage, ploidy and LD provenance','Signed LD and optional environment covariance; complete grid required','Simulations, LD stress and Arabidopsis case'],
        ['MRBIGR [17]','Population-scale multi-omics, GWAS and MR toolbox','Maize and rice case studies; broad multi-omics workflow','Different MR/network estimands; not the present G×E covariance contract','Literature scope comparison'],
        ['MR-Base / TwoSampleMR [10]','Human GWAS repository and two-sample MR automation','Human-oriented public GWAS ecosystem','Standard MR sensitivity analyses; plant metadata not the primary contract','Literature scope comparison'],
        ['metaGE [25]','Multi-environment GWAS meta-analysis','Plant METs and G×E QTL detection','Models heterogeneity and environmental covariates at GWAS meta-analysis level','Estimand comparison; not reimplemented'],
        ['MAPtools [36]','Mapping-by-sequencing and QTL-Seq command-line analysis','Plant-tested, multi-species mapping workflows','Variant/QTL mapping rather than summary-statistics MR covariance','Workflow/software-article comparison'],
    ],[2.6,4.1,4.2,4.8,3.3],7.2)

    heading(doc,'Environment-effect-heterogeneity model',2)
    para(doc,'For SNP j and environment k, the ratio estimate is r_jk = beta_y,jk / beta_x,jk. Let z_k be a prespecified numeric environmental score. PlantMR fits r_jk = theta_0 + theta_1 z_k + epsilon_jk. theta_0 is the genetically proxied effect at z=0 and theta_1 is the change in the MR effect per unit increase in z. This is an effect-heterogeneity model. Its intercept is not the MR-GxE pleiotropy intercept and is not interpreted as a correction for horizontal pleiotropy.')
    para(doc,'The first-order delta variance is v_jk = se_y,jk² / beta_x,jk² + beta_y,jk² se_x,jk² / beta_x,jk⁴. With SNP-major/environment-major ordering, optional signed SNP and environment correlation matrices define V = D(R_SNP ⊗ R_ENV)D, where D is the diagonal matrix of ratio standard errors. The GLS estimator is theta_hat = (X′V⁻¹X)⁻¹X′V⁻¹r, with X=[1,z]. If V is rank-deficient, residual heterogeneity degrees of freedom are rank(V)−rank(X).')
    para(doc,'The model requires a complete environment grid for each SNP and retains only SNPs that pass exposure P value, F statistic and MAF requirements in every environment. This prevents an apparent environmental slope from being caused solely by changing instrument composition. Missing environment correlation is not treated as verified independence; the report states that a diagonal approximation was used.')

    heading(doc,'Implementation details',2)
    para(doc,'The study was organized as a software/methods evaluation. We did not use the real-data case to select a favorable estimator, define a new gene-level claim or tune simulation truth. The main claims concern input validation, covariance-aware estimation and diagnostic transparency.')
    heading(doc,'Input schema and harmonization',3)
    para(doc,'Required summary columns were SNP, effect_allele, other_allele, beta, se and pval. Optional columns were eaf and n. The schema rejects empty/duplicate SNP identifiers, non-finite values, non-positive SE, invalid P values, invalid allele characters and equal effect/other alleles. Harmonization recognizes aligned, reversed, complemented and complemented-reversed alleles and uses allele frequencies to retain only resolvable palindromic variants.')
    heading(doc,'Instrument selection',3)
    para(doc,'The primary thresholds were exposure P≤5×10−8, F≥10 and MAF≥0.05 in the Arabidopsis case. The GxE selector requires every retained SNP to pass these filters in every environment. This prevents environment-specific tool composition from masquerading as effect heterogeneity. LD is represented by signed correlations, not r² values.')
    heading(doc,'Estimators and diagnostics',3)
    para(doc,'Standard methods include Wald ratio, fixed-effects IVW, random-effects IVW, MR-Egger, Cochran Q and leave-one-out analysis. The environment-slope estimator uses the delta-method ratio variance and GLS covariance described above. If the covariance is singular, a Moore–Penrose inverse is used for the quadratic form and Q degrees of freedom are rank adjusted.')
    heading(doc,'Simulation design',3)
    para(doc,'All simulations used fixed seeds recorded in metadata JSON files. The base benchmark used 20 SNPs and four environment scores (−1.5, −0.5, 0.5, 1.5). Exposure effects were generated as positive strong instruments. Outcome effects were generated from a pooled effect plus a known environment slope, with multivariate normal sampling errors. The LD stress benchmark additionally used an AR(1) LD correlation matrix with rho=0.6.')
    heading(doc,'Arabidopsis data processing',3)
    para(doc,'The GSE80744 normalized matrix was downloaded from GEO and the AT1G11560 row extracted. The 1001 Genomes HDF5 matrix was used to extract the local region; five PCs were calculated from every 1000th marker. A public accession VCF supplied REF/ALT labels. AraPheno FT10 and FT16 values were merged by accession ID. Local association regressions included genotype and PC1–PC5. Source URLs, hashes, generated tables and limitations are recorded with the case.')
    heading(doc,'Statistical reporting',3)
    para(doc,'All numerical results are reported with effect estimates, standard errors, P values, instrument counts, covariance source and heterogeneity diagnostics, following the reporting principles of STROBE-MR.[13] No genome-wide multiple-testing claim is made for the single-gene Arabidopsis demonstration. Any future genome-wide run must prespecify gene-level aggregation, multiple-testing control and an independent validation set.')

    heading(doc,'Results',1)
    heading(doc,'Simulation calibration and LD stress',2)
    para(doc,'The primary simulation benchmark used 20 instruments, four environments, environmental correlation 0.5, exposure SE 0.01, outcome SE 0.04 and 500 replicates per scenario. It included a null slope, a true slope of 0.25 and a common directional-pleiotropic component. The LD stress benchmark used the same dimensions with AR(1) signed LD correlation rho=0.6 and compared correct covariance specification with a diagonal independence approximation.')
    para(doc,'Table 3. Primary simulation calibration and LD-stress summary. Dashed lines in Figures 2 and 3 mark nominal 0.05 rejection and 0.95 coverage targets where applicable.',after=2)
    table(doc,['Scenario','Covariance','Mean/bias','Coverage','Rejection or power'],[
        ['Null','Environment-aware','slope bias −0.00069','0.946','0.054'],
        ['Null','Diagonal','slope bias −0.00071','0.994','0.006'],
        ['Causal G×E','Environment-aware','mean 0.24982; bias −0.00018','0.952','1.000'],
        ['Causal G×E','Diagonal','mean 0.24832; bias −0.00168','1.000','1.000'],
        ['LD stress null','Correct LD+environment','bias 0.00172','0.956','0.044'],
        ['LD stress null','Diagonal misspecified','bias 0.00177','0.874','0.126'],
        ['LD stress causal','Correct LD+environment','bias 0.00136','0.970','1.000'],
        ['LD stress causal','Diagonal misspecified','bias 0.00170','0.880','1.000'],
    ],[3.5,4.3,4.2,2.7,3.0],8.0)
    para(doc,'The directional-pleiotropy scenario exposed an important limitation. The environment slope remained approximately unbiased when the true slope was zero, but the pooled intercept was biased by +0.216. Therefore, a stable environment slope cannot be interpreted as proof that the pooled causal relation is protected from horizontal pleiotropy.')
    if FIG.exists():
        doc.add_picture(str(FIG),width=Inches(6.35))
        p=doc.add_paragraph('Figure 2. Primary simulation benchmark. The panels summarize slope estimation, rejection behavior and 95% confidence-interval coverage across null, causal environment-slope and directional-pleiotropy scenarios.')
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic=True; p.runs[0].font.size=Pt(9)

    if FIG_LD.exists():
        doc.add_picture(str(FIG_LD),width=Inches(6.25))
        p=doc.add_paragraph('Figure 3. LD stress benchmark. Treating correlated SNP and environment errors as independent preserved approximate point estimates but produced smaller standard errors, inflated null rejection from 0.044 to 0.126 and reduced coverage from 0.956 to 0.874.')
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic=True; p.runs[0].font.size=Pt(9)

    heading(doc,'Arabidopsis data-contract case',2)
    para(doc,'The Arabidopsis case was designed as a reproducibility and diagnostics demonstration rather than a new causal discovery. Baseline leaf expression of AT1G11560 was extracted from the GSE80744 normalized expression matrix, local genotypes were taken from the 1001 Genomes v3.1 matrix and flowering time at 10°C and 16°C was taken from AraPheno. The local association model used ordinary least squares with five whole-genome PCs, following the general principle that population structure and kinship must be modeled in plant association analyses,[29–34] but it was not presented as a reimplementation of the published mixed-model/SMR analysis.[19–24]')
    para(doc,'Table 4. Arabidopsis data-contract inputs and local processing.',after=2)
    table(doc,['Layer','Public source','Local processing'],[
        ['Molecular exposure','GSE80744 normalized expression','AT1G11560 row; log1p transformation; baseline exposure'],
        ['Genotype','1001 Genomes v3.1','Chr1:3,861,124–3,901,085; 3,352 local SNPs; five PCs'],
        ['Outcome environment 1','AraPheno FT10','10°C flowering time; z=−3'],
        ['Outcome environment 2','AraPheno FT16','16°C flowering time; z=+3'],
        ['Alleles','1001 Genomes accession VCF','REF/ALT mapping for dosage-coded matrix'],
    ],[4.0,5.0,8.0],8.2)
    para(doc,'At P≤5×10−8, F≥10 and MAF≥0.05, 48 SNPs passed in both environments. They were strongly correlated, so the primary analysis supplied a signed LD correlation matrix. The primary diagonal-environment-covariance run estimated an intercept of 2.3601 (SE 0.4685, P=4.72×10−7) and an environment slope of −0.03705 (SE 0.15617, P=0.81249). Residual heterogeneity was Q=192.89 on 60 rank-adjusted degrees of freedom (P=7.33×10−16).')
    para(doc,'Table 5. Instrument audit for the Arabidopsis case. Exclusion counts are overlapping diagnostics and must not be added as if they were mutually exclusive.',after=2)
    table(doc,['Audit stage','Count','Interpretation'],[
        ['Raw complete SNP×environment grid','172 SNPs / 344 rows','Two environments were available for every raw local SNP.'],
        ['Ambiguous palindromic exclusions','3 per environment','Removed during allele harmonization; 169 SNPs / 338 rows remained.'],
        ['P-value diagnostic exclusions','121 rows','Overlapping with other QC exclusions; not a unique-stage count.'],
        ['F-statistic diagnostic exclusions','106 rows','Overlapping with P-value and complete-grid diagnostics.'],
        ['Final complete-grid instruments','48 SNPs / 96 rows','Passed the prespecified P-value, F and MAF requirements in both environments.'],
        ['LD matrix / covariance rank','48×48 LD; rank(V)=62','Rank-aware Q used 62−2=60 residual degrees of freedom.'],
    ],[5.0,4.0,8.0],8.0)
    para(doc,'Table 6. Environment-stratified comparator estimates. Estimates are local OLS summary-statistics demonstrations and are not independent validation of the causal model.',after=2)
    table(doc,['Environment','Estimator','Estimate','SE','P value','Q (P value)'],[
        ['10°C','Fixed IVW','6.6001','0.2038','4.999×10−230','84.90 (5.89×10−4)'],
        ['10°C','Random IVW','7.0842','0.2970','1.030×10−125','84.90 (5.89×10−4)'],
        ['10°C','MR-Egger','4.2259','0.5188','3.758×10−16','—'],
        ['16°C','Fixed IVW','8.2379','0.2933','1.402×10−173','97.70 (2.05×10−5)'],
        ['16°C','Random IVW','8.8667','0.4561','3.598×10−84','97.70 (2.05×10−5)'],
        ['16°C','MR-Egger','6.7552','0.7683','1.466×10−18','—'],
    ],[2.2,3.0,3.0,2.3,3.5,4.0],7.6)
    para(doc,'The two environment-stratified IVW estimates are both positive but differ in magnitude (fixed IVW 6.6001 at 10°C versus 8.2379 at 16°C). Their heterogeneity statistics are significant, which motivates—but does not by itself identify—the environment-slope analysis. MR-Egger estimates are lower and have large intercept diagnostics (10°C intercept 2.9281, P=3.59×10−8; 16°C intercept 2.0160, P=0.0103), consistent with the need to treat pleiotropy as an unresolved limitation rather than a solved problem.')
    para(doc,'A sensitivity run used the Pearson correlation between FT10 and FT16 across 1,122 shared AraPheno accessions (r=0.88195) as a phenotype-correlation proxy. It estimated an intercept of 1.8027 (SE 0.6217, P=0.00374) and slope −0.08228 (SE 0.06634, P=0.21491). This proxy is not known ratio-error covariance and is not treated as primary inference.')
    note(doc,'Interpretation','The case shows that PlantMR can make plant data provenance, LD dependence, complete-grid selection and covariance approximations visible. It does not prove AT1G11560 causality, because expression and outcome cohorts share accessions, exposure–outcome covariance is not estimated, local OLS differs from the published LMM and residual heterogeneity remains.',YELLOW)

    if FIG_CASE.exists():
        doc.add_picture(str(FIG_CASE),width=Inches(6.55))
        p=doc.add_paragraph('Figure 4. Arabidopsis data-contract case. (A) Environment-stratified comparator estimates; (B) primary and phenotype-correlation-proxy environment slopes; (C) the raw-to-final instrument audit. Error bars are 95% confidence intervals where applicable.')
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic=True; p.runs[0].font.size=Pt(9)

    heading(doc,'Reproducibility and runtime',2)
    para(doc,'The local runtime benchmark completed validation, a synthetic G×E run and the Arabidopsis case with return code 0. The three tasks required 9.70–9.93 seconds and peak resident memory of 126,664–133,424 KB in the frozen environment. These are reproducibility benchmarks on one host, not claims of universal performance or genome-wide scalability.')
    para(doc,'Table 7. Local runtime benchmark for the frozen PlantMR implementation.',after=2)
    table(doc,['Task','Wall time (s)','Peak RSS (KB)','Return code'],[
        ['validate','9.7034','126,664','0'],
        ['gxe_synthetic','9.7974','129,736','0'],
        ['gxe_arabidopsis','9.9300','133,424','0'],
    ],[6.0,4.0,4.0,3.0],8.6)

    heading(doc,'Discussion',1)
    heading(doc,'Principal findings',2)
    para(doc,'This study makes three separable contributions. First, PlantMR turns plant context into an explicit computational contract rather than an unstructured note. Species, assembly, tissue, stage, environmental score, ploidy, allele coding, LD source and overlap assumptions travel with the summary statistics. This matters because the same SNP-level beta and standard error can represent different estimands when the expression tissue, developmental stage or treatment changes. A method that ignores those fields may still return a numerical estimate, but it cannot establish that two estimates are comparable.')
    para(doc,'Second, the environment-slope component provides a transparent estimand for a specific question: does the genetically proxied exposure–outcome effect change along a prespecified environmental scale? The slope is not a universal gene-by-environment causal parameter, and it is not a pleiotropy-correction intercept. Its value is that the design matrix, instrument grid and covariance assumptions are visible. The complete-grid rule is particularly important: if one environment has one set of instruments and another environment has a different set, a difference in ratio estimates can be produced by instrument composition rather than biological effect modification.')
    para(doc,'Third, the calibration results identify covariance specification as an inferential, not cosmetic, choice. Under the stated simulation truth, both correct and diagonal covariance models produced similar point estimates. The independence approximation nevertheless reduced standard errors and caused anti-conservative rejection in the LD stress null scenario (0.126 rather than 0.044), with coverage falling to 0.874. This pattern is precisely the type of failure that can be missed when a method is evaluated only by bias or point-estimate correlation. The results support a reporting rule: when SNP or environment dependence is known or plausible, the covariance source must be declared and a diagonal analysis should be presented as sensitivity, not as the default truth.')

    heading(doc,'Relationship to previous MR and plant methods',2)
    para(doc,'PlantMR should be interpreted alongside, not as a replacement for, established MR methods. IVW, MR-Egger, robust or median estimators, pleiotropy diagnostics, colocalization and fine-mapping address different threats and answer different questions.[3–12,28] The current implementation includes a limited set of standard estimators to provide a common audit trail, but it does not claim that one estimator solves invalid instruments. In particular, the large Q statistic and non-zero Egger intercepts in the Arabidopsis case show that the dataset contains unresolved heterogeneity and possible pleiotropy.')
    para(doc,'The distinction from MR-GxE is substantive. MR-GxE was developed to use gene-by-covariate interactions in the instrument–exposure association to detect or correct bias under assumptions about pleiotropy.[14] Interaction-based MR and MR-GENIUS use a different identification logic and have their own requirements for interaction strength and heterogeneity.[15] MR-EILLS targets an invariant causal effect across heterogeneous GWAS summaries and is not equivalent to fitting an environment score to two ratio estimates.[16] PlantMR instead supplies a plant-focused data contract and estimates effect heterogeneity along an observed environmental scale. It should not be marketed as a novel general interaction-MR theory or as an implementation of those methods.')
    para(doc,'The distinction from plant-oriented software is deliberate. MRBIGR integrates population-scale multi-omics analysis in maize and is a relevant comparator for software breadth.[17] PlantMR is narrower in biological scope but more explicit about plant summary-statistics metadata, signed LD, environmental covariance, complete SNP-by-environment grids and rank-aware Q degrees of freedom. Recent Plant Methods software articles such as MAPtools combine a workflow figure, implementation details, command-level outputs, published-data tests and an explicit availability block.[36] PlantMR follows that reader path: the workflow comes before the statistical details, the tool contract is separated from biological claims, and the real case is presented as a reproducibility demonstration. The remaining difference is important: MAPtools includes direct testing across published mapping datasets, whereas PlantMR currently provides a literature-scope matrix and controlled calibration rather than a formal head-to-head comparison against external MR implementations. This is a submission limitation, not evidence that PlantMR outperforms those tools.')

    heading(doc,'Why LD, environmental dependence and rank matter',2)
    para(doc,'Plant panels often contain related accessions, local haplotypes and uneven LD. The covariance of ratio estimates therefore cannot be assumed to be diagonal merely because the input file has one row per SNP. The same issue occurs across environments when the same accessions are phenotyped repeatedly, when trials share controls or when environmental conditions are derived from related measurements. Treating these observations as independent typically affects the standard errors and heterogeneity statistic more directly than the point estimate.')
    para(doc,'The rank-aware Q calculation is a practical safeguard for singular or nearly singular covariance matrices. With 48 retained SNPs and two environments, the case had 96 ratio observations but a covariance rank of 62. Using 94 or 95 degrees of freedom would overstate the independent information. PlantMR therefore reports rank(V)−rank(X), which was 60 in the primary case. This is not a guarantee that the covariance matrix is correctly specified; it is a guarantee that the reported degrees of freedom are at least aligned with the supplied matrix rank. A future implementation should add condition-number diagnostics, positive-semidefinite repair logs and sensitivity to alternative LD panels.')
    para(doc,'The stress benchmark also illustrates a useful distinction between calibration and power. Under a true slope of 0.25, both covariance choices had power 1.00, so a power-only comparison would have judged them equivalent. Under the null, the diagonal model failed calibration. For software papers, this argues for reporting null rejection, confidence coverage, bias, RMSE and standard-error behavior together, rather than reporting only the ability to detect a large simulated effect.')

    heading(doc,'Interpretation of the Arabidopsis case',2)
    para(doc,'The Arabidopsis case is intentionally a stress test of the data contract. The raw local region contained a complete two-environment grid, three ambiguous palindromic variants were removed per environment, and 48 SNPs passed the final complete-grid instrument filters. The environment-stratified fixed IVW estimates were positive at both temperatures but differed in magnitude, while heterogeneity Q statistics were significant. The covariance-aware slope estimate was negative but imprecise and not statistically distinguishable from zero. The phenotype-correlation proxy changed the point estimate and standard error, but it is not a measured covariance of the exposure and outcome GWAS estimates.')
    para(doc,'These findings do not support a claim that AT1G11560 is a validated causal regulator of flowering time. The expression and outcome accessions overlap, the exact exposure–outcome sampling covariance is unavailable, the local association model is ordinary least squares rather than a mixed model and the exposure was a baseline expression proxy copied across the two outcome environments rather than an environment-specific eQTL. The strong residual heterogeneity and Egger intercept diagnostics further limit a simple causal interpretation. The appropriate conclusion is that the software can expose these problems and produce reproducible sensitivity results; the biological conclusion remains provisional.')
    para(doc,'This conservative interpretation is a strength for a methods paper. A less cautious manuscript could quote the highly significant environment-specific IVW P values and call AT1G11560 causal, but that would ignore LD, pleiotropy, sample overlap, population structure and the difference between a local re-analysis and independent validation. PlantMR is designed to make that overstatement harder by placing warnings and provenance in the report itself.')

    heading(doc,'Implications for plant and crop research',2)
    para(doc,'For natural-accession studies, the data contract should record the accession panel, geographic structure strategy, kinship or mixed-model method, reference assembly and the allele representation used to construct the LD matrix. For breeding panels and multi-environment trials, the environmental score should be defined before analysis and should represent a biologically interpretable contrast rather than a post hoc label. Temperature, drought severity, photoperiod and nutrient level may be continuous, but their measurement error and correlation across trials should be documented.')
    para(doc,'For expression-mediated analyses, an eQTL association must be tied to tissue, developmental stage, treatment and expression normalization. A baseline eQTL should not be silently used to represent a stress-specific molecular exposure. When the biological question is mediation, colocalization, fine-mapping or HEIDI-type evidence is needed in addition to MR.[11,12,28] When the question is gene-by-environment interaction, the exposure association itself may need environment-specific modeling; the present environment-slope model should not be presented as a substitute.')
    para(doc,'For polyploid crops, dosage, homoeolog assignment and subgenome-specific LD are unresolved extensions rather than hidden defaults. Similarly, presence/absence variants and structural variants require explicit encoding and allele-frequency definitions. Keeping these features outside v1.1 is preferable to silently treating them as diploid SNPs. The boundary is documented so that users can distinguish a software limitation from a biological negative result.')

    heading(doc,'Limitations and threats to validity',2)
    numbered(doc,'The real case is not independent causal validation. It uses public processed data, overlapping accessions and a local OLS model. It demonstrates execution and diagnostics only.')
    numbered(doc,'The exact covariance of exposure and outcome summary statistics is not available. The environment-correlation sensitivity uses a phenotype correlation proxy and therefore cannot be treated as a formal sample-overlap correction.')
    numbered(doc,'Only two outcome environments are used in the case. A two-point slope is a contrast along a prespecified scale; it cannot establish linearity, a threshold, or a general response curve.')
    numbered(doc,'The directional-pleiotropy simulation uses one structured scenario. It does not cover balanced pleiotropy, correlated pleiotropy, weak instruments, nonlinear effects, measurement error in the environmental score, winner’s curse or population-stratified LD mismatch.')
    numbered(doc,'The current implementation accepts supplied signed correlation matrices but does not estimate a causal LD panel, repair every possible non-positive-semidefinite input or model polyploid dosage and structural variation.')
    numbered(doc,'Formal head-to-head implementations of MR-GxE, MR-GENIUS and MR-EILLS are not included in v1.1.0. Literature comparisons are conceptual and not benchmark claims.')
    numbered(doc,'The source package is reproducible locally but a public repository, DOI, clean-environment rerun and independent validation are still required for a final software publication record.')

    heading(doc,'Recommended validation ladder and future work',2)
    para(doc,'A credible next-stage biological application should proceed in layers. First, rerun the same exposure and outcome in non-overlapping accessions and estimate the exposure–outcome covariance under the actual GWAS design. Second, reproduce association statistics with a kinship-aware mixed model and use the matching LD panel. Third, compare PlantMR with formal colocalization/HEIDI and with a prespecified set of external interaction-MR methods under identical instruments, environments and outcome scales. Fourth, validate the direction and magnitude in an independent population or environment. Fifth, test a small number of genes with functional perturbation or expression validation before making pathway-level claims.')
    para(doc,'Software development should add condition-number and positive-semidefinite diagnostics, sample-overlap covariance interfaces, weak-instrument sensitivity, nonlinear environmental bases, multi-environment mixed-model importers, and explicit polyploid/SV schemas. Every extension should be accompanied by null calibration, misspecification stress tests and a negative-control scenario. The current v1.1.0 boundary is therefore a stable foundation for a broader platform, not a claim that the plant MR problem is solved.')

    heading(doc,'Conclusions',1)
    para(doc,'PlantMR is a reproducible plant-focused implementation for auditable summary-statistics MR and a narrowly defined environment-effect-heterogeneity analysis. Its strongest evidence is software-level: explicit provenance, complete-grid selection, covariance-aware GLS, rank-aware heterogeneity reporting, calibrated simulations and a real-data contract that preserves limitations. The evidence does not support calling AT1G11560 a causal gene or calling PlantMR the first plant environment-interaction MR method. A Plant Methods software submission is defensible after adding an archival DOI, clean-environment verification and completed author declarations; a biological discovery claim would require a substantially stronger validation program.')

    heading(doc,'Availability and requirements',1)
    para(doc,'PlantMR is released under the MIT License. The software snapshot is v1.1.0 and the complete manuscript/benchmark/Word-document package is frozen at tag v1.1.0-paper. The public repository is https://github.com/Zhangzhishuai-HIT/PlantMR, and the source archive `release/PlantMR_v1.1.0-paper_source.zip` is included for submission as supplementary software. A Zenodo or equivalent archival DOI remains to be assigned.')
    para(doc,'Public data sources include GSE80744 normalized expression data, AraPheno FT10/FT16, 1001 Genomes v3.1 and the Arabidopsis source publications. The large 1001 Genomes provider archive is not redistributed; the derived local genotype region, PC scores, allele mapping and provider checksums are included. The maize supplementary workbook, eQTL table and candidate table are included for audit but not used as a new causal result.')
    para(doc,'Table 8. Plant Methods software availability and requirements.',after=2)
    table(doc,['Field','Current value'],[
        ['Project name','PlantMR 1.1'],
        ['Project home page','https://github.com/Zhangzhishuai-HIT/PlantMR'],
        ['Operating system(s)','Linux verified; platform-independent Python code intended.'],
        ['Programming language','Python 3.10 or later.'],
        ['Dependencies','NumPy, pandas, SciPy, statsmodels, matplotlib and psutil; declared in environment.yml.'],
        ['License','MIT License.'],
        ['Restrictions for non-academic use','No additional restriction stated in the local license.'],
    ],[6.0,11.0],8.4)

    heading(doc,'List of abbreviations',1)
    table(doc,['Abbreviation','Definition'],[
        ['MR','Mendelian randomization'],
        ['GWAS','Genome-wide association study'],
        ['eQTL','Expression quantitative trait locus'],
        ['G×E / GxE','Genotype-by-environment interaction or environment-effect heterogeneity'],
        ['LD','Linkage disequilibrium'],
        ['GLS','Generalized least squares'],
        ['IVW','Inverse-variance weighted'],
        ['MAF','Minor allele frequency'],
        ['SNP','Single-nucleotide polymorphism'],
        ['Q','Cochran-type heterogeneity statistic'],
    ],[4.5,12.5],8.4)

    heading(doc,'Declarations',1)
    table(doc,['Declaration','Status'],[
        ['Ethics approval and consent to participate','Not applicable: public plant accessions and public aggregate/processed data were used.'],
        ['Consent for publication','Not applicable.'],
        ['Availability of data and materials','Public source data and derived audit materials are described above; code is available at https://github.com/Zhangzhishuai-HIT/PlantMR; archival DOI pending.'],
        ['Competing interests','To be completed by authors.'],
        ['Funding','To be completed by authors.'],
        ['Author contributions','To be completed by authors.'],
        ['Acknowledgements','To be completed by authors or marked Not applicable.'],
        ['Authors’ information','Author names, affiliations and corresponding author to be supplied.'],
    ],[4.5,12.5],8.8)

    heading(doc,'References',1)
    refs=[
        '1. Davey Smith, G. & Ebrahim, S. “Mendelian randomization”: can genetic epidemiology contribute to understanding environmental determinants of disease? Int. J. Epidemiol. 32, 1–22 (2003). https://doi.org/10.1093/ije/dyg070',
        '2. Lawlor, D. A., Harbord, R. M., Sterne, J. A. C., Timpson, N. & Smith, G. D. Mendelian randomization: using genes as instruments for making causal inferences in epidemiology. Stat. Med. 27, 1133–1163 (2008). https://doi.org/10.1002/sim.3034',
        '3. Burgess, S., Butterworth, A. & Thompson, S. G. Mendelian randomization analysis with multiple genetic variants using summarized data. Genet. Epidemiol. 37, 658–665 (2013). https://doi.org/10.1002/gepi.21758',
        '4. Bowden, J., Davey Smith, G. & Burgess, S. Mendelian randomization with invalid instruments: effect estimation and bias detection through Egger regression. Int. J. Epidemiol. 44, 512–525 (2015). https://doi.org/10.1093/ije/dyv080',
        '5. Bowden, J., Davey Smith, G., Haycock, P. C. & Burgess, S. Consistent estimation in Mendelian randomization with some invalid instruments using a weighted median estimator. Genet. Epidemiol. 40, 304–314 (2016). https://doi.org/10.1002/gepi.21965',
        '6. Hartwig, F. P., Davey Smith, G. & Bowden, J. Robust inference in summary data Mendelian randomization via the zero modal pleiotropy assumption. Int. J. Epidemiol. 46, 1985–1998 (2017). https://doi.org/10.1093/ije/dyx102',
        '7. Verbanck, M. et al. Detection of widespread horizontal pleiotropy in causal relationships inferred from Mendelian randomization between complex traits and diseases. Nat. Genet. 50, 693–698 (2018). https://doi.org/10.1038/s41588-018-0099-7',
        '8. Zhao, Q., Wang, J., Hemani, G., Bowden, J. & Small, D. S. Statistical inference in two-sample summary-data Mendelian randomization using robust adjusted profile score. Ann. Statist. 48, 1742–1769 (2020). https://doi.org/10.1214/19-AOS1866',
        '9. Burgess, S. et al. Guidelines for performing Mendelian randomization investigations: update for summer 2023. Wellcome Open Res. 4, 186 (2023). https://doi.org/10.12688/wellcomeopenres.15555.3',
        '10. Hemani, G. et al. The MR-Base platform supports systematic causal inference across the human phenome. eLife 7, e34408 (2018). https://doi.org/10.7554/eLife.34408',
        '11. Zhu, Z. et al. Integration of summary data from GWAS and eQTL studies predicts complex trait gene targets. Nat. Genet. 48, 481–487 (2016). https://doi.org/10.1038/ng.3538',
        '12. Zhu, Z. et al. Causal associations between risk factors and common diseases inferred from GWAS summary data. Nat. Commun. 9, 224 (2018). https://doi.org/10.1038/s41467-017-02317-2',
        '13. Skrivankova, V. W. et al. Strengthening the reporting of observational studies in epidemiology using Mendelian randomization: the STROBE-MR statement. JAMA 326, 1614–1621 (2021). https://doi.org/10.1001/jama.2021.18236',
        '14. Spiller, W. et al. Detecting and correcting for bias in Mendelian randomization analyses using Gene-by-Environment interactions. Int. J. Epidemiol. 48, 702–712 (2019). https://doi.org/10.1093/ije/dyy204',
        '15. Spiller, W., Hartwig, F. P., Sanderson, E., Davey Smith, G. & Bowden, J. Interaction-based Mendelian randomization with measured and unmeasured gene-by-covariate interactions. PLoS ONE 17, e0271933 (2022). https://doi.org/10.1371/journal.pone.0271933',
        '16. Hou, L., Chen, H. & Zhou, X.-H. MR-EILLS: an invariance-based Mendelian randomization method integrating multiple heterogeneous GWAS summary datasets. Nat. Commun. 16, 1–15 (2025). https://doi.org/10.1038/s41467-025-62823-6',
        '17. Xu, F. et al. MRBIGR: a versatile toolbox for genetic regulation inference from population-scale multi-omics data. Plant Commun. 6, 101197 (2025). https://doi.org/10.1016/j.xplc.2024.101197',
        '18. Liu, S. et al. Mapping regulatory variants controlling gene expression in drought response and tolerance in maize. Genome Biol. 21, 163 (2020). https://doi.org/10.1186/s13059-020-02069-1',
        '19. Feng, X. et al. Dual-trait genomic analysis in highly stratified Arabidopsis thaliana populations using genome-wide association summary statistics. Heredity 133, 11–20 (2024). https://doi.org/10.1038/s41437-024-00688-z',
        '20. The 1001 Genomes Consortium. 1,135 Genomes reveal the global pattern of polymorphism in Arabidopsis thaliana. Cell 166, 481–491 (2016). https://doi.org/10.1016/j.cell.2016.05.063',
        '21. Kawakatsu, T. et al. Epigenomic diversity in a global collection of Arabidopsis thaliana accessions. Cell 166, 492–505 (2016). https://doi.org/10.1016/j.cell.2016.06.044',
        '22. Schmitz, R. J. et al. Patterns of population epigenomic diversity. Nature 495, 193–198 (2013). https://doi.org/10.1038/nature11968',
        '23. Atwell, S. et al. Genome-wide association study of 107 phenotypes in Arabidopsis thaliana inbred lines. Nature 465, 627–631 (2010). https://doi.org/10.1038/nature08800',
        '24. Brachi, B. et al. Linkage and association mapping of Arabidopsis thaliana flowering time in nature. PLoS Genet. 6, e1000940 (2010). https://doi.org/10.1371/journal.pgen.1000940',
        '25. De Walsche, A. et al. metaGE: investigating genotype × environment interactions through GWAS meta-analysis. PLoS Genet. 21, e1011553 (2025). https://doi.org/10.1371/journal.pgen.1011553',
        '26. Sul, J. H. et al. Accounting for population structure in gene-by-environment interactions in genome-wide association studies using mixed models. PLoS Genet. 12, e1005849 (2016). https://doi.org/10.1371/journal.pgen.1005849',
        '27. An, X. et al. An approach to identify gene-environment interactions and reveal new biological insight in complex traits. Nat. Commun. 15, 1–16 (2024). https://doi.org/10.1038/s41467-024-47806-3',
        '28. Giambartolomei, C. et al. Bayesian test for colocalisation between pairs of genetic association studies using summary statistics. PLoS Genet. 10, e1004383 (2014). https://doi.org/10.1371/journal.pgen.1004383',
        '29. Price, A. L. et al. Principal components analysis corrects for stratification in genome-wide association studies. Nat. Genet. 38, 904–909 (2006). https://doi.org/10.1038/ng1847',
        '30. Kang, H. M. et al. Efficient control of population structure in model organism association mapping. Genetics 178, 1709–1723 (2008). https://doi.org/10.1534/genetics.107.080101',
        '31. Zhou, X. & Stephens, M. Genome-wide efficient mixed-model analysis for association studies. Nat. Genet. 44, 821–824 (2012). https://doi.org/10.1038/ng.2310',
        '32. Lipka, A. E. et al. GAPIT: genome association and prediction integrated tool. Bioinformatics 28, 2397–2399 (2012). https://doi.org/10.1093/bioinformatics/bts444',
        '33. Purcell, S. et al. PLINK: a tool set for whole-genome association and population-based linkage analyses. Am. J. Hum. Genet. 81, 559–575 (2007). https://doi.org/10.1086/519795',
        '34. Chang, C. C. et al. Second-generation PLINK: rising to the challenge of larger and richer datasets. GigaScience 4, 7 (2015). https://doi.org/10.1186/s13742-015-0047-8',
        '35. Kang, M. et al. The pan-genome and local adaptation of Arabidopsis thaliana. Nat. Commun. 14, 6259 (2023). https://doi.org/10.1038/s41467-023-42029-4',
        '36. Candela, H. et al. MAPtools: command-line tools for mapping-by-sequencing and QTL-Seq analysis and visualization. Plant Methods 20, 107 (2024). https://doi.org/10.1186/s13007-024-01222-2',
    ]
    for ref in refs:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(0.45); p.paragraph_format.first_line_indent=Cm(-0.45); p.paragraph_format.space_after=Pt(3); p.add_run(ref).font.size=Pt(8.5)

    cp=doc.core_properties; cp.title='PlantMR: plant and crop Mendelian randomization with environment-aware effect heterogeneity'; cp.subject='Submission manuscript'; cp.author='PlantMR contributors'; cp.keywords='Mendelian randomization, plants, crops, LD, GxE, software'
    OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT)

if __name__=='__main__': main()
